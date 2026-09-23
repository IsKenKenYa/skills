
import sys
from pathlib import Path
import logging
import random
from typing import Optional

import pandas as pd


from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Side, Border, PatternFill
from openpyxl.utils import range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parent / "../_lib"))
from encoding import *  # noqa: E402  # Force UTF-8 output (fixes Windows GBK garbled text)
sys.path.insert(0, str(Path(__file__).resolve().parent / "."))
sys.path.insert(0, str(Path(__file__).resolve().parent / "so_field"))
from config import config_data, TraceType, MAX_INT
from callchain import OptimizedCallChain
from hidumper import read_showmap_2_df


logger = logging.getLogger()


HIPROFILER_NOISE_NAME = "anon_inode:dev/ashmem/unique_stack_table"


class TopdownContext:
    __slots__ = (
        "hidumper_total_size",
        "arkts_heap_total_size",
        "arkts_heap_decompose_size",
        "arkts_heap_uncover",
        "arkts_heap_empty_service_size",
        "arkts_heap_remain",
        "native_heap_total_size",
        "native_heap_decompose_size",
        "native_heap_empty_service_size",
        "native_heap_remain",
        "native_heap_brk",
        "native_heap_jemalloc_meta",
        "native_heap_meta",
        "native_heap_mmap",
        "native_profile_noise",
        "gl_total_size",
        "gl_decompose_size",
        "gl_empty_service_size",
        "gl_remain",
        "guard_total_size",
        "guard_decompose_size",
        "guard_empty_service_size",
        "guard_remain",
        "graph_total_size",
        "graph_decompose_size",
        "graph_empty_service_size",
        "graph_remain",
        "graph_pixelmap_size",
        "graph_pixelmap_ratio",
        "hap_total_size",
        "anon_page_total_size",
        "anon_page_decompose_size_p1",
        "anon_page_empty_service_size",
        "anon_page_remain",
        "db_total_size",
        "db_decompose_size",
        "db_empty_service_size",
        "db_remain",
        "so_total_size",
        "so_decompose_size",
        "dev_total_size",
        "dev_decompose_size",
        "dev_empty_service_size",
        "dev_remain",
        "ttf_total_size",
        "ttf_decompose_size",
        "ttf_empty_service_size",
        "ttf_remain",
        "file_page_total_size",
        "file_page_decompose_size",
        "file_page_noise_size",
        "file_page_empty_service_size",
        "file_page_remain",
        "file_page_profile_noise",
    )

    def __init__(self):
        self.hidumper_total_size = 0
        self.arkts_heap_total_size = 0
        self.arkts_heap_decompose_size = 0
        self.arkts_heap_uncover = 0
        self.arkts_heap_empty_service_size = 0
        self.arkts_heap_remain = 0
        self.native_heap_total_size = 0
        self.native_heap_empty_service_size = 0
        self.native_heap_decompose_size = 0
        self.native_heap_brk = 0
        self.native_heap_jemalloc_meta = 0
        self.native_heap_meta = 0
        self.native_heap_mmap = 0
        self.native_heap_remain = 0
        self.native_profile_noise = 0
        self.gl_total_size = 0
        self.gl_decompose_size = 0
        self.gl_empty_service_size = 0
        self.gl_remain = 0
        self.guard_total_size = 0
        self.guard_decompose_size = 0
        self.guard_empty_service_size = 0
        self.guard_remain = 0
        self.graph_total_size = 0
        self.graph_decompose_size = 0
        self.graph_empty_service_size = 0
        self.graph_remain = 0
        self.graph_pixelmap_size = 0
        self.graph_pixelmap_ratio = 0
        self.hap_total_size = 0
        self.anon_page_total_size = 0
        self.anon_page_decompose_size_p1 = 0
        self.anon_page_empty_service_size = 0
        self.anon_page_remain = 0
        self.db_total_size = 0
        self.db_decompose_size = 0
        self.db_empty_service_size = 0
        self.db_remain = 0
        self.so_total_size = 0
        self.so_decompose_size = 0
        self.dev_total_size = 0
        self.dev_decompose_size = 0
        self.dev_empty_service_size = 0
        self.dev_remain = 0
        self.ttf_total_size = 0
        self.ttf_decompose_size = 0
        self.ttf_empty_service_size = 0
        self.ttf_remain = 0
        self.file_page_total_size = 0
        self.file_page_decompose_size = 0
        self.file_page_noise_size = 0
        self.file_page_empty_service_size = 0
        self.file_page_remain = 0
        self.file_page_profile_noise = 0


# noinspection DuplicatedCode,PyBroadException
def topdown_analyze(callchains: OptimizedCallChain) -> Optional[TopdownContext]:
    context = TopdownContext()
    try:
        native_heap_decompose_size = 0
        arkts_heap_decompose_size = 0
        gl_decompose_size = 0
        graph_decompose_size = 0
        file_page_decompose_size_p1 = 0
        native_heap_brk = 0
        native_heap_jemalloc_meta = 0
        native_heap_meta = 0
        native_heap_mmap = 0
        anon_page_decompose_size = 0
        for cc in callchains:
            if cc["type"] == TraceType.NATIVE_HEAP.name:
                if cc["field"]["response_so"] == "[anon:native_heap:brk]":
                    native_heap_brk += cc["heap_size"]
                elif cc["field"]["response_so"] == "[anon:native_heap:jemalloc meta]":
                    native_heap_jemalloc_meta += cc["heap_size"]
                elif cc["field"]["response_so"] == "[anon:native_heap:meta]":
                    native_heap_meta += cc["heap_size"]
                elif cc["field"]["response_so"] == "[anon:native_heap:mmap]":
                    native_heap_mmap += cc["heap_size"]
                else:
                    native_heap_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.ARKTS_HEAP.name:
                arkts_heap_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.GPU_VK.name:
                gl_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.GPU_GLES.name:
                gl_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.GPU_CL.name:
                gl_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.DMA.name:
                graph_decompose_size += cc["heap_size"]
            elif cc["type"] == TraceType.ASHMEM.name:
                file_page_decompose_size_p1 += cc["heap_size"]
            elif cc["type"] == TraceType.ANON_PAGE_OTHER.name:
                anon_page_decompose_size += cc["heap_size"]
        native_heap_decompose_size /= (1024 * 1024)
        arkts_heap_decompose_size /= (1024 * 1024)
        gl_decompose_size /= (1024 * 1024)
        graph_decompose_size /= (1024 * 1024)
        file_page_decompose_size_p1 /= (1024 * 1024)
        native_heap_brk /= (1024 * 1024)
        native_heap_jemalloc_meta /= (1024 * 1024)
        native_heap_meta /= (1024 * 1024)
        native_heap_mmap /= (1024 * 1024)
        anon_page_decompose_size /= (1024 * 1024)
        # 1. mmap 拆解出来的 category: size
        proc_smaps_file = Path(config_data["output_path"]) / f"混合模式明细-{config_data['scene']}.xlsx"
        if proc_smaps_file.exists():
            smaps_df = pd.read_excel(proc_smaps_file)
            category_sum_series = smaps_df.groupby('Category')['Calculated_Rss'].sum()
            decompose_size = category_sum_series.to_dict()
        else:
            decompose_size = {}

        # 2. meminfo.xlsx 中的各 Category 总内存（单位 MB）
        df = pd.read_excel(Path(config_data["output_path"]) / "meminfo.xlsx")
        # 3. trace 中拆出来的 native heap、arkts heap、dma、gpu 的 size
        # 4. Filepage Other 的 hiprofiler 底噪
        showmap_df = read_showmap_2_df(Path(config_data["showmap"]))
        # Handle showmap formats that use "Pss" instead of "Rss"
        if "Rss" not in showmap_df.columns and "Pss" in showmap_df.columns:
            showmap_df["Rss"] = pd.to_numeric(showmap_df["Pss"], errors="coerce").fillna(0)
        elif "Rss" in showmap_df.columns:
            showmap_df["Rss"] = pd.to_numeric(showmap_df["Rss"], errors="coerce").fillna(0)
        else:
            showmap_df["Rss"] = 0
        # Handle showmap without "Name" column (no hiprofiler noise filtering)
        if "Name" in showmap_df.columns:
            filtered_df = showmap_df[showmap_df['Name'] == HIPROFILER_NOISE_NAME]
            file_page_noise_size = filtered_df['Rss'].sum() / 1024
        else:
            file_page_noise_size = 0

        mask = df["time"] == config_data.get("scene", config_data["endts"])
        total_size = df.loc[mask].drop(columns=["time", "stack"]).sum(axis=1).iloc[0]
        arkts_heap_total_size = (
            df.loc[mask, "ark ts heap"].iloc[0]
            if "ark ts heap" in df.columns and not df.loc[mask].empty
            else 0
        )
        native_heap_total_size = (
            df.loc[mask, "native heap"].iloc[0]
            if "native heap" in df.columns and not df.loc[mask].empty
            else 0
        )
        native_rows = showmap_df[showmap_df['Category'] == 'native heap']
        native_heap_empty_service_size = config_data.get("empty_service", 0)
        native_profile_noise = config_data.get("hiprofiler_noise", 0)
        # Cap native_heap_decompose_size so that breakdown doesn't exceed total
        native_decompose_budget = max(
            native_heap_total_size - native_profile_noise - native_heap_empty_service_size
            - native_heap_brk - native_heap_jemalloc_meta - native_heap_meta - native_heap_mmap,
            0
        )
        native_heap_decompose_size = min(native_heap_decompose_size, native_decompose_budget)
        native_heap_remain = max(
            native_heap_total_size - native_heap_decompose_size - native_profile_noise - native_heap_empty_service_size - native_heap_brk
            - native_heap_jemalloc_meta - native_heap_meta - native_heap_mmap,
            0
        )
        arkts_heap_empty_service_size = config_data.get("arkts_empty_service", 0)
        # Cap arkts_heap_decompose_size so that breakdown doesn't exceed total
        arkts_decompose_budget = max(arkts_heap_total_size - arkts_heap_empty_service_size, 0)
        arkts_heap_decompose_size = min(arkts_heap_decompose_size, arkts_decompose_budget)
        arkts_heap_uncover = max((arkts_heap_decompose_size / 0.9) * 0.1 - arkts_heap_empty_service_size, 0)
        # Recompute budget after uncover
        arkts_decompose_budget = max(arkts_heap_total_size - arkts_heap_uncover - arkts_heap_empty_service_size, 0)
        arkts_heap_decompose_size = min(arkts_heap_decompose_size, arkts_decompose_budget)
        arkts_heap_remain = max(
            arkts_heap_total_size - arkts_heap_decompose_size - arkts_heap_uncover - arkts_heap_empty_service_size,
            0
        )
        gl_total_size = df.loc[mask].filter(regex=r"^gpu").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        gl_empty_service_size = 0
        gl_remain = max(gl_total_size - gl_decompose_size - gl_empty_service_size, 0)
        guard_total_size = (
            df.loc[mask].filter(regex=r"^guard").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        )
        guard_empty_service_size = (
            df.loc[mask, "guard(empty_service)"].iloc[0] if "guard(empty_service)" in df.columns else 0
        )
        guard_decompose_size = decompose_size.get("guard", 0) / 1024
        guard_remain = max(guard_total_size - guard_decompose_size - guard_empty_service_size, 0)
        graph_total_size = df.loc[mask].filter(regex=r"^dma").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        graph_empty_service_size = 0
        graph_remain = max(graph_total_size - graph_decompose_size - graph_empty_service_size, 0)
        mm_dmabuf_file = config_data.get("mm_dmabuf_file", "")
        graph_pixelmap_size = 0
        graph_pixelmap_ratio = 0
        if mm_dmabuf_file and Path(mm_dmabuf_file).exists():
            _dma_df = _parse_mm_dmabuf(mm_dmabuf_file)
            if _dma_df is not None and not _dma_df.empty:
                graph_pixelmap_size = round(_dma_df[_dma_df["buf_type"] == "pixelmap"]["size_mb"].sum(), 2)
                graph_pixelmap_ratio = round(graph_pixelmap_size / graph_total_size * 100, 2) if graph_total_size > 0 else 0
        hap_total_size = df.loc[mask].filter(regex=r"^\.hap").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        hap_empty_service_size = (
            df.loc[mask, ".hap(empty_service)"].iloc[0] if ".hap(empty_service)" in df.columns else 0
        )
        hap_decompose_size = decompose_size.get(".hap", 0) / 1024
        hap_remain = max(hap_total_size - hap_decompose_size - hap_empty_service_size, 0)
        anon_page_total_size = (
            df.loc[mask].filter(regex=r"^AnonPage other").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        )
        anon_empty_name = "AnonPage other(empty_service)"
        anon_page_empty_service_size = (
            df.loc[mask, anon_empty_name].iloc[0] if anon_empty_name in df.columns else 0
        )
        anon_page_decompose_size = min(anon_page_decompose_size, max(anon_page_total_size - anon_page_empty_service_size, 0))
        anon_page_remain = max(anon_page_total_size - anon_page_decompose_size - anon_page_empty_service_size, 0)
        stack_total_size = (
            df.loc[mask].filter(regex=r"^stack").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        )
        stack_empty_service_size = (
            df.loc[mask, "stack(empty_service)"].iloc[0] if "stack(empty_service)" in df.columns else 0
        )
        stack_decompose_size = decompose_size.get("stack", 0) / 1024
        stack_remain = max(stack_total_size - stack_decompose_size - stack_empty_service_size, 0)
        db_total_size = df.loc[mask].filter(regex=r"^\.db").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        db_empty_service_size = (
            df.loc[mask, ".db(empty_service)"].iloc[0] if ".db(empty_service)" in df.columns else 0
        )
        db_decompose_size = decompose_size.get(".db", 0) / 1024
        db_remain = max(db_total_size - db_decompose_size - db_empty_service_size, 0)
        so_total_size = df.loc[mask].filter(regex=r"^\.so").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        so_empty_service_size = (
            df.loc[mask, ".so(empty_service)"].iloc[0] if ".so(empty_service)" in df.columns else 0
        )
        so_decompose_size = so_total_size
        dev_total_size = df.loc[mask].filter(regex=r"^dev").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        dev_empty_service_size = (
            df.loc[mask, "dev(empty_service)"].iloc[0] if "dev(empty_service)" in df.columns else 0
        )
        dev_decompose_size = decompose_size.get("dev", 0) / 1024
        dev_remain = max(dev_total_size - dev_decompose_size - dev_empty_service_size, 0)
        ttf_total_size = df.loc[mask].filter(regex=r"^\.ttf").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        ttf_empty_service_size = (
            df.loc[mask, ".ttf(empty_service)"].iloc[0] if ".ttf(empty_service)" in df.columns else 0
        )
        ttf_decompose_size = decompose_size.get(".ttf", 0) / 1024
        ttf_remain = max(ttf_total_size - ttf_decompose_size - ttf_empty_service_size, 0)
        file_page_total_size = (
            df.loc[mask].filter(regex=r"^FilePage other").apply(pd.to_numeric, errors="coerce").sum(axis=1).iloc[0]
        )
        file_page_empty_service_size = (
            df.loc[mask, "FilePage other(empty_service)"]
            .iloc[0] if "FilePage other(empty_service)" in df.columns else 0
        )
        file_page_decompose_size = (decompose_size.get("FilePage other", 0) / 1024) + file_page_decompose_size_p1
        file_page_remain = max(
            file_page_total_size - file_page_decompose_size - file_page_empty_service_size - file_page_noise_size, 0
        )
        context.hidumper_total_size = total_size
        context.arkts_heap_total_size = arkts_heap_total_size
        context.arkts_heap_decompose_size = arkts_heap_decompose_size
        context.arkts_heap_uncover = arkts_heap_uncover
        context.arkts_heap_empty_service_size = arkts_heap_empty_service_size
        context.arkts_heap_remain = arkts_heap_remain
        context.native_heap_total_size = native_heap_total_size
        context.native_heap_decompose_size = native_heap_decompose_size
        context.native_heap_empty_service_size = native_heap_empty_service_size
        context.native_heap_brk = native_heap_brk
        context.native_heap_jemalloc_meta = native_heap_jemalloc_meta
        context.native_heap_meta = native_heap_meta
        context.native_heap_mmap = native_heap_mmap
        context.native_heap_remain = native_heap_remain
        context.native_profile_noise = native_profile_noise
        context.gl_total_size = gl_total_size
        context.gl_decompose_size = gl_decompose_size
        context.gl_empty_service_size = gl_empty_service_size
        context.gl_remain = gl_remain
        context.guard_total_size = guard_total_size
        context.guard_decompose_size = guard_decompose_size
        context.guard_empty_service_size = guard_empty_service_size
        context.guard_remain = guard_remain
        context.graph_total_size = graph_total_size
        context.graph_decompose_size = graph_decompose_size
        context.graph_empty_service_size = graph_empty_service_size
        context.graph_remain = graph_remain
        context.graph_pixelmap_size = graph_pixelmap_size
        context.graph_pixelmap_ratio = graph_pixelmap_ratio
        context.hap_total_size = hap_total_size
        context.anon_page_total_size = anon_page_total_size
        context.anon_page_decompose_size_p1 = anon_page_decompose_size
        context.anon_page_empty_service_size = anon_page_empty_service_size
        context.anon_page_remain = anon_page_remain
        context.db_total_size = db_total_size
        context.db_decompose_size = db_decompose_size
        context.db_empty_service_size = db_empty_service_size
        context.db_remain = db_remain
        context.so_total_size = so_total_size
        context.so_decompose_size = so_decompose_size
        context.dev_total_size = dev_total_size
        context.dev_decompose_size = dev_decompose_size
        context.dev_empty_service_size = dev_empty_service_size
        context.dev_remain = dev_remain
        context.ttf_total_size = ttf_total_size
        context.ttf_decompose_size = ttf_decompose_size
        context.ttf_empty_service_size = ttf_empty_service_size
        context.ttf_remain = ttf_remain
        context.file_page_total_size = file_page_total_size
        context.file_page_decompose_size = file_page_decompose_size
        context.file_page_noise_size = file_page_noise_size
        context.file_page_empty_service_size = file_page_empty_service_size
        context.file_page_remain = file_page_remain
        context.file_page_profile_noise = file_page_noise_size

        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        center = Alignment(horizontal="center", vertical="center")
        bold = Font(bold=True)
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="FFFF00")

        ws.merge_cells("A1:AM1")
        ws["A1"] = f"页面总内存 {round(total_size, 2)} MB"
        ws["A1"].font = bold
        ws["A1"].alignment = center
        ws["A1"].fill = header_fill

        ws["A2"] = f"arkts heap ({round(arkts_heap_total_size, 2)} MB)"
        ws["E2"] = f"native heap ({round(native_heap_total_size, 2)} MB)"
        ws["M2"] = f"GL ({round(gl_total_size, 2)} MB)"  # 原I2
        ws["P2"] = f"guard ({round(guard_total_size, 2)} MB)"  # 原L2
        ws["S2"] = f"Graph ({round(graph_total_size, 2)} MB)"  # 原O2
        ws["T2"] = f"pixelmap占比 ({graph_pixelmap_ratio}%)"
        ws["W2"] = f"hap ({round(hap_total_size, 2)} MB)"  # 原R2
        ws["X2"] = f"AnonPage other ({round(anon_page_total_size, 2)} MB)"  # 原S2
        ws["AA2"] = f".db ({round(db_total_size, 2)} MB)"  # 原V2
        ws["AD2"] = f".so ({round(so_total_size, 2)} MB)"  # 原Y2
        ws["AE2"] = f"dev ({round(dev_total_size, 2)} MB)"  # 原AB2
        ws["AH2"] = f".ttf ({round(ttf_total_size, 2)} MB)"  # 原AE2
        ws["AJ2"] = f"FilePage other ({round(file_page_total_size, 2)} MB)"  # 原AH2
        # arkts heap部分
        ws["A3"] = "可拆解内存"
        ws["A4"] = f"{round(arkts_heap_decompose_size, 2)} MB"
        ws["B3"] = "未覆盖"
        ws["B4"] = f"{round(arkts_heap_uncover, 2)} MB"
        ws["C3"] = "空服务"
        ws["C4"] = f"{round(arkts_heap_empty_service_size, 2)} MB"
        ws["D3"] = "分配器碎片"
        ws["D4"] = f"{round(arkts_heap_remain, 2)} MB"
        # native heap部分
        ws["E3"] = "可拆解内存"  # 原E3
        ws["E4"] = f"{round(native_heap_decompose_size, 2)} MB"  # 原E4
        ws["F3"] = "native_heap:brk"
        ws["F4"] = f"{round(native_heap_brk, 2)} MB"
        ws["G3"] = "native_heap:jemalloc meta"
        ws["G4"] = f"{round(native_heap_jemalloc_meta, 2)} MB"
        ws["H3"] = "native_heap:meta"
        ws["H4"] = f"{round(native_heap_meta, 2)} MB"
        ws["I3"] = "native_heap:mmap"
        ws["I4"] = f"{round(native_heap_mmap, 2)} MB"
        ws["J3"] = "hiprofiler底噪"  # 原F3
        ws["J4"] = f"{round(native_profile_noise, 2)} MB"
        ws["K3"] = "空服务"  # 原G3
        ws["K4"] = f"{round(native_heap_empty_service_size, 2)} MB"  # 原G4
        ws["L3"] = "缓存&分配器碎片&未覆盖"  # 原H3
        ws["L4"] = f"{round(native_heap_remain, 2)} MB"  # 原H4
        # GL部分
        ws["M3"] = "可拆解内存"  # 原I3
        ws["M4"] = f"{round(gl_decompose_size, 2)} MB"  # 原I4
        ws["N3"] = "空服务"  # 原J3
        ws["N4"] = f"{round(gl_empty_service_size, 2)} MB"  # 原J4
        ws["O3"] = "分配器缓存&元数据&4K对齐"  # 原K3
        ws["O4"] = f"{round(gl_remain, 2)} MB"  # 原K4

        # guard部分
        ws["P3"] = "可拆解内存"  # 原L3
        ws["P4"] = f"{round(guard_decompose_size, 2)} MB"  # 原L4
        ws["Q3"] = "空服务"  # 原M3
        ws["Q4"] = f"{round(guard_empty_service_size, 2)} MB"  # 原M4
        ws["R3"] = "未覆盖"  # 原N3
        ws["R4"] = f"{round(guard_remain, 2)} MB"  # 原N4

        # graph部分
        ws["S3"] = "可拆解内存"  # 原O3
        ws["S4"] = f"{round(graph_decompose_size, 2)} MB"  # 原O4
        ws["T3"] = "pixelmap"  # 新增
        ws["T4"] = f"{graph_pixelmap_size} MB ({graph_pixelmap_ratio}%)"  # 新增
        ws["U3"] = "空服务"  # 原P3
        ws["U4"] = f"{round(graph_empty_service_size, 2)} MB"  # 原P4
        ws["V3"] = "未覆盖"  # 原Q3
        ws["V4"] = f"{round(graph_remain, 2)} MB"  # 原Q4

        # DMA明细（从mm_dmabuf_info解析，追加在topdown报告下方）
        mm_dmabuf_file = config_data.get("mm_dmabuf_file", "")
        if mm_dmabuf_file and Path(mm_dmabuf_file).exists():
            _append_dma_detail_to_topdown(ws, mm_dmabuf_file, graph_total_size)

        # hap部分
        ws["W3"] = "可拆解内存"  # 原R3
        ws["W4"] = f"{round(hap_total_size, 2)} MB"  # 原R4

        # AnonPage other部分
        ws["X3"] = "可拆解内存"  # 原S3
        ws["X4"] = f"{round(anon_page_decompose_size, 2)} MB"  # 原S4
        ws["Y3"] = "空服务"  # 原T3
        ws["Y4"] = f"{round(anon_page_empty_service_size, 2)} MB"  # 原T4
        ws["Z3"] = "未覆盖"  # 原U3
        ws["Z4"] = f"{round(anon_page_remain, 2)} MB"  # 原U4

        # .db部分
        ws["AA3"] = "可拆解内存"  # 原V3
        ws["AA4"] = f"{round(db_decompose_size, 2)} MB"  # 原V4
        ws["AB3"] = "空服务"  # 原W3
        ws["AB4"] = f"{round(db_empty_service_size, 2)} MB"  # 原W4
        ws["AC3"] = "未覆盖"  # 原X3
        ws["AC4"] = f"{round(db_remain, 2)} MB"  # 原X4

        # .so部分
        ws["AD3"] = "可拆解内存"  # 原Y3
        ws["AD4"] = f"{round(so_decompose_size, 2)} MB"  # 原Y4

        # dev部分
        ws["AE3"] = "可拆解内存"  # 原AB3
        ws["AE4"] = f"{round(dev_decompose_size, 2)} MB"  # 原AB4
        ws["AF3"] = "空服务"  # 原AC3
        ws["AF4"] = f"{round(dev_empty_service_size, 2)} MB"  # 原AC4
        ws["AG3"] = "未覆盖"  # 原AD3
        ws["AG4"] = f"{round(dev_remain, 2)} MB"  # 原AD4

        # .ttf部分
        ws["AH3"] = "可拆解内存"  # 原AE3
        ws["AH4"] = f"{round(ttf_decompose_size + ttf_remain, 2)} MB"  # 原AE4
        ws["AI3"] = "空服务"  # 原AF3
        ws["AI4"] = f"{round(ttf_empty_service_size, 2)} MB"  # 原AF4

        # FilePage other部分
        ws["AJ3"] = "可拆解内存"  # 原AH3
        ws["AJ4"] = f"{round(file_page_decompose_size + file_page_remain, 2)} MB"  # 原AH4
        ws["AK3"] = "hiprofiler底噪"  # 原AI3
        ws["AK4"] = f"{round(file_page_noise_size, 2)} MB"  # 原AI4
        ws["AL3"] = "空服务"  # 原AJ3
        ws["AL4"] = f"{round(file_page_empty_service_size, 2)} MB"  # 原AJ4
        ws["AN1"] = "总覆盖率"
        analiysisable = arkts_heap_total_size + native_heap_total_size + gl_total_size + guard_total_size + graph_total_size + \
            hap_total_size + anon_page_total_size + db_total_size + so_total_size + dev_total_size + \
            ttf_total_size + file_page_total_size - arkts_heap_uncover - native_heap_remain - guard_remain - \
            graph_remain - hap_remain - anon_page_remain - db_remain - dev_remain
        ws["AN2"] = f"{round(analiysisable/total_size * 100, 2)} %"
        merge_ranges = [
            "A2:D2", "E2:L2", "M2:O2", "P2:R2", "S2:V2",
            "W2:W2", "X2:Z2", "AA2:AC2", "AD2:AD2", "AE2:AG2",
            "AH2:AI2", "AJ2:AL2"
        ]
        for rng in merge_ranges:
            ws.merge_cells(rng)
            min_col, min_row, max_col, max_row = range_boundaries(rng)
            for row in ws.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col
            ):
                for cell in row:
                    cell.border = border
                    cell.font = bold
                    cell.alignment = center
                    cell.fill = header_fill
        for row in ws["A3:AM4"]:
            for cell in row:
                cell.alignment = center
        wb.save(Path(config_data["output_path"]) / f"topdown报告-{config_data['scene']}.xlsx")
    except Exception as e:
        logger.exception("生成 topdown 报告错误")
    return context

def _scale_callchain_heap_sizes(callchains: OptimizedCallChain, context: TopdownContext):
    """Scale heap_size of overcounted types so they don't exceed showmap total.

    htrace statistics mode can overcount (apply_size - release_size > actual physical memory)
    because cumulative counters lose precision at 1-second granularity.
    Scale proportionally so decompose_size matches the capped value from topdown.
    """
    _NATIVE_HEAP_EXCLUDED = (
        '[anon:native_heap:brk]', '[anon:native_heap:jemalloc meta]',
        '[anon:native_heap:meta]', '[anon:native_heap:mmap]'
    )

    # Collect raw totals before scaling
    raw_arkts = sum(cc['heap_size'] for cc in callchains if cc['type'] == TraceType.ARKTS_HEAP.name)
    raw_native = sum(
        cc['heap_size'] for cc in callchains
        if cc['type'] == TraceType.NATIVE_HEAP.name
        and cc['field'].get('response_so') not in _NATIVE_HEAP_EXCLUDED
    )
    raw_anon = sum(cc['heap_size'] for cc in callchains if cc['type'] == TraceType.ANON_PAGE_OTHER.name)

    # Target capped values (in bytes) from topdown context
    target_arkts = context.arkts_heap_decompose_size * 1024 * 1024
    target_native = context.native_heap_decompose_size * 1024 * 1024
    target_anon = context.anon_page_decompose_size_p1 * 1024 * 1024

    # Scale factors
    arkts_scale = (target_arkts / raw_arkts) if raw_arkts > 0 else 1.0
    native_scale = (target_native / raw_native) if raw_native > 0 else 1.0
    anon_scale = (target_anon / raw_anon) if raw_anon > 0 else 1.0

    for cc in callchains:
        if cc['type'] == TraceType.ARKTS_HEAP.name and arkts_scale < 1.0:
            cc['heap_size'] = cc['heap_size'] * arkts_scale
        elif (cc['type'] == TraceType.NATIVE_HEAP.name
              and cc['field'].get('response_so') not in _NATIVE_HEAP_EXCLUDED
              and native_scale < 1.0):
            cc['heap_size'] = cc['heap_size'] * native_scale
        elif cc['type'] == TraceType.ANON_PAGE_OTHER.name and anon_scale < 1.0:
            cc['heap_size'] = cc['heap_size'] * anon_scale


def add_topdown_data(callchains: OptimizedCallChain):
    context = topdown_analyze(callchains)
    if context:
        _scale_callchain_heap_sizes(callchains, context)
        callchains.append({
            'heap_size': context.arkts_heap_uncover * 1024 * 1024,
            'type': TraceType.ARKTS_HEAP.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "ArkTS Heap未覆盖",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.arkts_heap_empty_service_size * 1024 * 1024,
            'type': TraceType.ARKTS_HEAP.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "ArkTS Heap空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.arkts_heap_remain * 1024 * 1024,
            'type': TraceType.ARKTS_HEAP.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "ArkTS Heap分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.native_heap_empty_service_size * 1024 * 1024,
            'type': TraceType.NATIVE_HEAP.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "Native Heap空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.native_heap_remain * 1024 * 1024,
            'type': TraceType.NATIVE_HEAP.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "Native Heap分配器碎片&未覆盖",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.gl_empty_service_size * 1024 * 1024,
            'type': TraceType.GPU_CL.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "GL空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.gl_remain * 1024 * 1024,
            'type': TraceType.GPU_CL.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "GL分配器缓存",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.guard_empty_service_size * 1024 * 1024,
            'type': TraceType.GUARD.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "guard空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.guard_remain * 1024 * 1024,
            'type': TraceType.GUARD.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "guard分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.graph_empty_service_size * 1024 * 1024,
            'type': TraceType.DMA.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "DMA空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.graph_remain * 1024 * 1024,
            'type': TraceType.DMA.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "DMA分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.anon_page_empty_service_size * 1024 * 1024,
            'type': TraceType.ANON_PAGE_OTHER.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "AnonPage Other空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.anon_page_remain * 1024 * 1024,
            'type': TraceType.ANON_PAGE_OTHER.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "AnonPage Other分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.db_empty_service_size * 1024 * 1024,
            'type': TraceType.DB.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": ".db空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.db_remain * 1024 * 1024,
            'type': TraceType.DB.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": ".db分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.dev_empty_service_size * 1024 * 1024,
            'type': TraceType.DEV.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "dev空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.dev_remain * 1024 * 1024,
            'type': TraceType.DEV.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "dev分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.file_page_empty_service_size * 1024 * 1024,
            'type': TraceType.FILE_PAGE_OTHER.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "Filepage Other空服务",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
        callchains.append({
            'heap_size': context.file_page_remain * 1024 * 1024,
            'type': TraceType.FILE_PAGE_OTHER.name,
            'callchain_id': -1 * random.randint(1, MAX_INT),
            'field': {
                "response_so": "Filepage Other分配器碎片",
                "firstkind": '语言运行时',
                "secondkind": '分配器碎片&空服务',
                "thirdkind": '分配器碎片&空服务'
            },
            'frames': [],

        })
    return context


def _parse_mm_dmabuf(file_path: str) -> Optional[pd.DataFrame]:
    """Parse mm_dmabuf_info file into a DataFrame with size_mb column."""
    try:
        lines = Path(file_path).read_text(encoding="utf-8", errors="ignore").splitlines()
        if len(lines) < 2:
            return None
        header = lines[0].split()
        data = [l.split() for l in lines[1:] if l.strip()]
        if not data:
            return None
        df = pd.DataFrame(data, columns=header)
        df["size_bytes"] = pd.to_numeric(df.get("size_bytes"), errors="coerce").fillna(0)
        df["size_mb"] = df["size_bytes"] / 1024 / 1024
        return df
    except Exception:
        return None


def _append_dma_detail_to_topdown(ws, mm_dmabuf_file: str, graph_total_size: float):
    """Append DMA buf_type and top-N image breakdown below the topdown table."""
    df = _parse_mm_dmabuf(mm_dmabuf_file)
    if df is None or df.empty:
        return

    start_row = 6
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    # Section: DMA by buf_type
    ws.cell(row=start_row, column=1, value="DMA buf_type 明细").font = bold
    ws.cell(row=start_row, column=1).fill = header_fill
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    start_row += 1

    by_type = df.groupby("buf_type")["size_mb"].agg(["sum", "count"]).sort_values("sum", ascending=False)
    ws.cell(row=start_row, column=1, value="buf_type").font = bold
    ws.cell(row=start_row, column=2, value="size (MB)").font = bold
    ws.cell(row=start_row, column=3, value="count").font = bold
    start_row += 1
    for buf_type, row in by_type.iterrows():
        ws.cell(row=start_row, column=1, value=buf_type)
        ws.cell(row=start_row, column=2, value=round(row["sum"], 2))
        ws.cell(row=start_row, column=3, value=int(row["count"]))
        start_row += 1
    ws.cell(row=start_row, column=1, value="合计").font = bold
    ws.cell(row=start_row, column=2, value=round(df["size_mb"].sum(), 2)).font = bold
    ws.cell(row=start_row, column=3, value=len(df)).font = bold
    start_row += 2

    # Section: DMA top 20 by image name
    ws.cell(row=start_row, column=1, value="DMA Top20 图片/资源明细").font = bold
    ws.cell(row=start_row, column=1).fill = header_fill
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=4)
    start_row += 1

    by_name = df.groupby("buf_name")["size_mb"].agg(["sum", "count"]).sort_values("sum", ascending=False).head(20)
    ws.cell(row=start_row, column=1, value="buf_name").font = bold
    ws.cell(row=start_row, column=2, value="size (MB)").font = bold
    ws.cell(row=start_row, column=3, value="count").font = bold
    ws.cell(row=start_row, column=4, value="buf_type").font = bold
    start_row += 1
    for buf_name, row in by_name.iterrows():
        name_display = buf_name[:60] if len(str(buf_name)) > 60 else buf_name
        ws.cell(row=start_row, column=1, value=name_display)
        ws.cell(row=start_row, column=2, value=round(row["sum"], 2))
        ws.cell(row=start_row, column=3, value=int(row["count"]))
        # Get dominant buf_type for this name
        dominant_type = df[df["buf_name"] == buf_name].groupby("buf_type")["size_bytes"].sum().idxmax()
        ws.cell(row=start_row, column=4, value=dominant_type)
        start_row += 1


def generate_dma_detail_excel(mm_dmabuf_file: str, output_path: Path):
    """Generate a standalone DMA detail Excel with buf_type and per-image breakdown."""
    df = _parse_mm_dmabuf(mm_dmabuf_file)
    if df is None or df.empty:
        return

    wb = Workbook()
    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)

    # Sheet 1: by buf_type
    ws1 = wb.active
    ws1.title = "DMA_by_buf_type"
    by_type = df.groupby("buf_type")["size_mb"].agg(["sum", "count"]).sort_values("sum", ascending=False).reset_index()
    by_type.columns = ["buf_type", "size_MB", "count"]
    by_type["percentage"] = (by_type["size_MB"] / by_type["size_MB"].sum() * 100).round(2)
    for r_idx, row in enumerate([by_type.columns.tolist()] + by_type.values.tolist(), 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = center
            if r_idx == 1:
                cell.font = bold

    # Sheet 2: by image name
    ws2 = wb.create_sheet("DMA_by_image")
    by_name = df.groupby("buf_name")["size_mb"].agg(["sum", "count"]).sort_values("sum", ascending=False).reset_index()
    by_name.columns = ["buf_name", "size_MB", "count"]
    by_name["percentage"] = (by_name["size_MB"] / by_name["size_MB"].sum() * 100).round(2)
    # Add dominant buf_type
    buf_type_map = df.groupby("buf_name")["buf_type"].apply(
        lambda x: x.value_counts().idxmax() if len(x) > 0 else ""
    ).to_dict()
    by_name["buf_type"] = by_name["buf_name"].map(buf_type_map)
    for r_idx, row in enumerate([by_name.columns.tolist()] + by_name.values.tolist(), 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = center
            if r_idx == 1:
                cell.font = bold

    # Sheet 3: raw data
    ws3 = wb.create_sheet("DMA_raw")
    raw_cols = ["buf_name", "size_bytes", "size_mb", "buf_type", "can_reclaim", "leak_type"]
    available_cols = [c for c in raw_cols if c in df.columns]
    raw_df = df[available_cols].sort_values("size_bytes", ascending=False)
    for r_idx, row in enumerate([raw_df.columns.tolist()] + raw_df.values.tolist(), 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = center
            if r_idx == 1:
                cell.font = bold

    wb.save(output_path / "dma_detail.xlsx")
